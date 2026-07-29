import os
import zipfile
from copy import deepcopy

import fitz
from docx import Document


def load_text_file(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as file:
        return file.read()


def load_docx_file(file_path: str) -> str:
    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"Uploaded DOCX was not found: {os.path.abspath(file_path)}"
        )

    if not zipfile.is_zipfile(file_path):
        raise ValueError(
            "The uploaded file has a .docx extension but is not a valid Word document."
        )

    document = Document(file_path)
    sections: list[str] = []

    # Normal document paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    # Include text inside tables
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                sections.append(" | ".join(cells))

    extracted_text = "\n\n".join(sections).strip()

    if not extracted_text:
        raise ValueError(
            "The uploaded Word document does not contain readable text."
        )

    return extracted_text


def load_pdf_file(file_path: str) -> str:
    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"Uploaded PDF was not found: {os.path.abspath(file_path)}"
        )

    pages: list[str] = []

    with fitz.open(file_path) as document:
        for page in document:
            text = page.get_text("text").strip()

            if text:
                pages.append(text)

    extracted_text = "\n\n".join(pages).strip()

    if not extracted_text:
        raise ValueError(
            "No readable text was found in this PDF. "
            "Scanned/image-only PDFs are not supported yet."
        )

    return extracted_text


def load_document(file_path: str) -> str:
    lower_path = file_path.lower()

    if lower_path.endswith((".txt", ".srt")):
        return load_text_file(file_path)

    if lower_path.endswith(".docx"):
        return load_docx_file(file_path)

    if lower_path.endswith(".pdf"):
        return load_pdf_file(file_path)

    raise ValueError(
        f"Unsupported document type: {os.path.splitext(file_path)[1]}"
    )


def copy_paragraph_format(source_paragraph, target_paragraph) -> None:
    """
    Copy paragraph-level formatting such as style, alignment,
    indentation, spacing and bullet/numbering properties.
    """
    if source_paragraph.style:
        target_paragraph.style = source_paragraph.style

    target_paragraph.alignment = source_paragraph.alignment

    source_format = source_paragraph.paragraph_format
    target_format = target_paragraph.paragraph_format

    target_format.left_indent = source_format.left_indent
    target_format.right_indent = source_format.right_indent
    target_format.first_line_indent = source_format.first_line_indent
    target_format.space_before = source_format.space_before
    target_format.space_after = source_format.space_after
    target_format.line_spacing = source_format.line_spacing
    target_format.keep_together = source_format.keep_together
    target_format.keep_with_next = source_format.keep_with_next
    target_format.page_break_before = source_format.page_break_before
    target_format.widow_control = source_format.widow_control

    # Copy low-level paragraph properties, including bullets/numbering.
    if source_paragraph._p.pPr is not None:
        target_paragraph._p.insert(
            0,
            deepcopy(source_paragraph._p.pPr),
        )


def copy_run_format(source_run, target_run) -> None:
    """
    Copy basic character-level formatting.
    """
    target_run.bold = source_run.bold
    target_run.italic = source_run.italic
    target_run.underline = source_run.underline

    target_run.font.name = source_run.font.name
    target_run.font.size = source_run.font.size
    target_run.font.bold = source_run.font.bold
    target_run.font.italic = source_run.font.italic
    target_run.font.underline = source_run.font.underline
    target_run.font.strike = source_run.font.strike
    target_run.font.subscript = source_run.font.subscript
    target_run.font.superscript = source_run.font.superscript

    if source_run.font.color is not None:
        try:
            target_run.font.color.rgb = source_run.font.color.rgb
        except (ValueError, TypeError):
            pass
