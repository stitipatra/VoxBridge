import os
from datetime import datetime
from pathlib import Path

from docx import Document
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

from app.services.document_service import (
    copy_paragraph_format,
    copy_run_format,
)

TRANSLATION_DIR = os.path.join("storage", "translations")

os.makedirs(TRANSLATION_DIR, exist_ok=True)


def save_text_output(text: str, output_type: str = "translation") -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    filename = f"{output_type}_{timestamp}.txt"
    file_path = os.path.join(TRANSLATION_DIR, filename)

    with open(file_path, "w", encoding="utf-8") as file:
        file.write(text)

    return file_path


def build_output_path(
    input_path: str,
    output_directory: str,
    extension: str,
) -> str:
    os.makedirs(output_directory, exist_ok=True)

    input_name = Path(input_path).stem

    return os.path.join(
        output_directory,
        f"{input_name}_translated{extension}",
    )


def save_docx_output(
    translated_text: str,
    input_path: str,
    output_directory: str = "storage/outputs/text",
) -> str:
    output_path = build_output_path(
        input_path=input_path,
        output_directory=output_directory,
        extension=".docx",
    )

    document = Document()

    translated_paragraphs = [
        paragraph.strip()
        for paragraph in translated_text.split("\n\n")
    ]

    for translated_paragraph in translated_paragraphs:
        if not translated_paragraph:
            document.add_paragraph()
            continue

        document.add_paragraph(translated_paragraph)

    document.save(output_path)

    return output_path


def save_formatted_docx_output(
    translated_paragraphs: list[str],
    input_path: str,
    output_directory: str = "storage/outputs/text",
) -> str:
    """
    Create a translated DOCX while preserving practical formatting from
    the original document.

    Preserves:
    - paragraph styles
    - headings
    - bullets and numbering
    - alignment
    - indentation and spacing
    - basic run styling
    - tables

    Complex layouts, floating elements and exact run-to-word formatting
    cannot always be preserved after translation.
    """
    source_document = Document(input_path)
    output_document = Document()

    # Remove the blank paragraph that a new Document may begin with.
    if (
        len(output_document.paragraphs) == 1
        and not output_document.paragraphs[0].text
    ):
        paragraph_element = output_document.paragraphs[0]._element
        paragraph_element.getparent().remove(paragraph_element)

    translation_index = 0

    for source_paragraph in source_document.paragraphs:
        target_paragraph = output_document.add_paragraph()

        copy_paragraph_format(
            source_paragraph,
            target_paragraph,
        )

        original_text = source_paragraph.text.strip()

        if not original_text:
            continue

        if translation_index >= len(translated_paragraphs):
            translated_text = original_text
        else:
            translated_text = translated_paragraphs[translation_index]

        translation_index += 1

        target_run = target_paragraph.add_run(translated_text)

        # Use the first source run as the practical formatting template.
        if source_paragraph.runs:
            copy_run_format(
                source_paragraph.runs[0],
                target_run,
            )

    # Preserve table structure and basic cell paragraph formatting.
    for source_table in source_document.tables:
        target_table = output_document.add_table(
            rows=len(source_table.rows),
            cols=len(source_table.columns),
        )

        try:
            target_table.style = source_table.style
        except (KeyError, ValueError):
            pass

        for row_index, source_row in enumerate(source_table.rows):
            for column_index, source_cell in enumerate(source_row.cells):
                target_cell = target_table.cell(
                    row_index,
                    column_index,
                )

                # Clear automatically created cell text.
                target_cell.text = ""

                for paragraph_index, source_paragraph in enumerate(
                    source_cell.paragraphs
                ):
                    if paragraph_index == 0:
                        target_paragraph = target_cell.paragraphs[0]
                    else:
                        target_paragraph = target_cell.add_paragraph()

                    copy_paragraph_format(
                        source_paragraph,
                        target_paragraph,
                    )

                    original_text = source_paragraph.text.strip()

                    if not original_text:
                        continue

                    if translation_index >= len(translated_paragraphs):
                        translated_text = original_text
                    else:
                        translated_text = translated_paragraphs[
                            translation_index
                        ]

                    translation_index += 1

                    target_run = target_paragraph.add_run(
                        translated_text
                    )

                    if source_paragraph.runs:
                        copy_run_format(
                            source_paragraph.runs[0],
                            target_run,
                        )

    # Copy section dimensions and margins.
    for index, source_section in enumerate(source_document.sections):
        if index >= len(output_document.sections):
            break

        target_section = output_document.sections[index]

        target_section.page_width = source_section.page_width
        target_section.page_height = source_section.page_height
        target_section.top_margin = source_section.top_margin
        target_section.bottom_margin = source_section.bottom_margin
        target_section.left_margin = source_section.left_margin
        target_section.right_margin = source_section.right_margin

    output_path = build_output_path(
        input_path=input_path,
        output_directory=output_directory,
        extension=".docx",
    )

    output_document.save(output_path)

    return output_path


def register_pdf_font() -> str:
    """
    Try to load a font that supports English and Devanagari.

    Windows generally contains Nirmala UI, which supports Hindi and
    Marathi. Falls back to Helvetica when unavailable.
    """
    possible_fonts = [
        r"C:\Windows\Fonts\Nirmala.ttf",
        r"C:\Windows\Fonts\NirmalaB.ttf",
        r"C:\Windows\Fonts\mangal.ttf",
    ]

    for font_path in possible_fonts:
        if os.path.exists(font_path):
            font_name = "AnuwadiniUnicode"

            try:
                pdfmetrics.registerFont(
                    TTFont(font_name, font_path)
                )
                return font_name
            except Exception:
                continue

    return "Helvetica"


def get_pdf_alignment(alignment_name: str | None) -> int:
    alignment_map = {
        "left": TA_LEFT,
        "center": TA_CENTER,
        "right": TA_RIGHT,
        "justify": TA_JUSTIFY,
    }

    return alignment_map.get(
        (alignment_name or "left").lower(),
        TA_LEFT,
    )


def escape_reportlab_text(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\n", "<br/>")
    )


def save_pdf_output(
    translated_text: str,
    input_path: str,
    output_directory: str = "storage/outputs/text",
) -> str:
    output_path = build_output_path(
        input_path=input_path,
        output_directory=output_directory,
        extension=".pdf",
    )

    font_name = register_pdf_font()
    styles = getSampleStyleSheet()

    body_style = ParagraphStyle(
        name="AnuwadiniBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=11,
        leading=17,
        alignment=TA_LEFT,
        spaceAfter=9,
    )

    document = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="Translated Document",
        author="Anuwadini",
    )

    story = []

    paragraphs = translated_text.split("\n\n")

    for paragraph_text in paragraphs:
        paragraph_text = paragraph_text.strip()

        if not paragraph_text:
            story.append(Spacer(1, 6))
            continue

        safe_text = escape_reportlab_text(paragraph_text)

        story.append(
            Paragraph(
                safe_text,
                body_style,
            )
        )

    document.build(story)

    return output_path
